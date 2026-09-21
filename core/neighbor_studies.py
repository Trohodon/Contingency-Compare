"""N-1 neighbor study processing through PowerWorld SimAuto."""

from __future__ import annotations

import csv
import tempfile
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


COMPANIES = (
    ("CPLE", "CPLE"),
    ("DUK", "DUKE"),
    ("SC", "Santee Cooper"),
    ("SOCO", "Southern Company"),
)
RESULTS_DIR = "Neighbor Study Result Files"
SOURCE_SUFFIX = "_CA_NOT_RUN_N-1.PWB"
WORKING_SUFFIX = "_ACCA_N-1.PWB"
EXPORT_FIELDS = ("CTGLabel", "LimViolID", "LimViolLimit", "LimViolValue", "LimViolPct", "LimViolCat")


@dataclass(frozen=True)
class StudyCase:
    study: str
    source: Path
    working: Path


def discover_n1_cases(root: Path) -> tuple[list[StudyCase], list[str]]:
    """Find one N-1 source case per immediate study folder."""
    cases: list[StudyCase] = []
    warnings: list[str] = []
    for study_dir in sorted(root.iterdir(), key=lambda p: p.name.lower()):
        if not study_dir.is_dir() or study_dir.name == RESULTS_DIR:
            continue
        n1_dir = next((p for p in study_dir.iterdir() if p.is_dir() and p.name.lower() == "n-1"), None)
        if n1_dir is None:
            warnings.append(f"{study_dir.name}: no N-1 folder")
            continue
        matches = sorted(
            (p for p in n1_dir.iterdir() if p.is_file() and p.name.upper().endswith(SOURCE_SUFFIX)),
            key=lambda p: p.name.lower(),
        )
        if len(matches) != 1:
            warnings.append(f"{study_dir.name}: expected one *{SOURCE_SUFFIX} case; found {len(matches)}")
            continue
        source = matches[0]
        working = source.with_name(source.name[:-len(SOURCE_SUFFIX)] + WORKING_SUFFIX)
        cases.append(StudyCase(study_dir.name, source, working))
    return cases, warnings


def _check(result, action: str) -> None:
    error = result[0] if isinstance(result, (tuple, list)) else result
    if error:
        raise RuntimeError(f"{action}: {error}")


def _script_path(path: Path) -> str:
    return str(path).replace("\\", "/").replace('"', '""')


def _read_export(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.reader(stream)
        first = next(reader, [])
        if len(first) == 1 and first[0].strip().lower() == "violationctg":
            headers = next(reader, [])
        else:
            headers = first
        if not headers:
            return []
        missing = set(EXPORT_FIELDS) - set(headers)
        if missing:
            raise RuntimeError(f"PowerWorld export is missing fields: {', '.join(sorted(missing))}")
        return [dict(zip(headers, row)) for row in reader if any(value.strip() for value in row)]


def _export_company(simauto, study: StudyCase, code: str, name: str, output_dir: Path, temp_dir: Path, log) -> list[dict[str, str]]:
    filter_name = f"{code}_P1"
    csv_path = temp_dir / f"{code}.csv"
    fields = ",".join(EXPORT_FIELDS)
    _check(simauto.RunScriptCommand(
        f'SaveData("{_script_path(csv_path)}", CSV, ViolationCTG, [{fields}], [], "{filter_name}");'
    ), f"Export {filter_name} violations")
    rows = _read_export(csv_path)
    labels = sorted({row["CTGLabel"].strip() for row in rows if row["CTGLabel"].strip()})
    log(f"  {name}: {len(rows)} violation rows, {len(labels)} contingencies")

    # Reset selection for each company, including when the preceding company had no rows.
    _check(simauto.RunScriptCommand("SetData(Contingency, [Selected], [NO], ALL);"), "Clear contingency selection")
    con_path = output_dir / f"{study.study}_{name} P1.con"
    if not labels:
        if con_path.exists():
            con_path.unlink()
            log(f"  Removed previous CON file with no matching P1 violations: {con_path}")
        return rows
    for label in labels:
        _check(simauto.ChangeParametersSingleElement(
            "Contingency", ["CTGLabel", "Selected"], [label, "YES"]
        ), f"Select contingency {label}")

    output_dir.mkdir(parents=True, exist_ok=True)
    # Stage alongside the destination so replacing a result also works on network drives.
    temp_con_path = output_dir / f".{uuid4().hex}.con"
    # The PTI exporter applies SELECTED to Contingency, unlike the result filter above.
    try:
        _check(simauto.RunScriptCommand(
            f'CTGWriteFilePTI("{_script_path(temp_con_path)}", Number, NO, "SELECTED", NO);'
        ), f"Write {name} PTI contingency file")
        if not temp_con_path.exists() or temp_con_path.stat().st_size == 0:
            raise RuntimeError(f"PowerWorld did not write a nonempty CON file: {temp_con_path}")
        temp_con_path.replace(con_path)
    finally:
        temp_con_path.unlink(missing_ok=True)
    log(f"  Wrote {con_path}")
    return rows


def run_n1_case(study: StudyCase, assets_dir: Path, results_root: Path, log, simauto_factory=None) -> dict[str, list[dict[str, str]]]:
    """Save a working case, apply settings, solve, and export four P1 groups."""
    if simauto_factory is None:
        import win32com.client
        simauto_factory = lambda: win32com.client.Dispatch("pwrworld.SimulatorAuto")
    settings = assets_dir / "NeighborStudy_N-1ContsSettings.aux"
    if not settings.is_file():
        raise FileNotFoundError(settings)
    simauto = simauto_factory()
    try:
        _check(simauto.OpenCase(str(study.source)), f"Open {study.source}")
        _check(simauto.SaveCase(str(study.working), "PWB", True), "Save ACCA working copy")
        _check(simauto.ProcessAuxFile(str(settings)), "Apply N-1 settings")
        _check(simauto.RunScriptCommand("EnterMode(Contingency);"), "Enter contingency mode")
        log(f"{study.study}: solving N-1 contingencies; this may take a while")
        _check(simauto.RunScriptCommand("CTGSolveAll;"), "Solve all contingencies")
        _check(simauto.SaveCase(str(study.working), "PWB", True), "Save solved ACCA case")
        log(f"{study.study}: saved solved case {study.working}")
        results = {}
        with tempfile.TemporaryDirectory(prefix="neighbor_study_") as folder:
            temp_dir = Path(folder)
            for code, name in COMPANIES:
                results[code] = _export_company(
                    simauto, study, code, name, results_root / name, temp_dir, log
                )
        _check(simauto.RunScriptCommand("SetData(Contingency, [Selected], [NO], ALL);"), "Clear contingency selection")
        return results
    finally:
        try:
            simauto.CloseCase()
        except Exception:
            pass


def _sheet_name(name: str, used: set[str]) -> str:
    cleaned = "".join("_" if ch in '[]:*?/\\' else ch for ch in name).strip() or "Study"
    base = cleaned[:31]
    candidate = base
    number = 2
    while candidate.lower() in used:
        suffix = f" ({number})"
        candidate = base[:31-len(suffix)] + suffix
        number += 1
    used.add(candidate.lower())
    return candidate


def write_company_workbooks(results_root: Path, studies: dict[str, dict[str, list[dict[str, str]]]], log) -> list[Path]:
    """One company workbook with one formatted sheet per completed study."""
    written = []
    navy = PatternFill("solid", fgColor="305496")
    white_bold = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="B7C9DF")
    headers = ("Contingency Events", "Resulting Issue", "Limit", "Contingency Value", "Percent Loading", "Category")
    fields = EXPORT_FIELDS
    for code, name in COMPANIES:
        company_dir = results_root / name
        company_dir.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        book.remove(book.active)
        used: set[str] = set()
        for study_name, company_rows in studies.items():
            ws = book.create_sheet(_sheet_name(study_name, used))
            ws.sheet_view.showGridLines = False
            ws.merge_cells("B2:G2")
            title = ws["B2"]
            title.value = f"{study_name} | {name} P1"
            title.fill = navy
            title.font = Font(color="FFFFFF", bold=True, size=12)
            title.alignment = Alignment(horizontal="center")
            for col, header in enumerate(headers, 2):
                cell = ws.cell(3, col, header)
                cell.fill = navy
                cell.font = white_bold
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            rows = company_rows.get(code, [])
            for row_number, item in enumerate(rows, 4):
                for col, field in enumerate(fields, 2):
                    raw = item.get(field, "")
                    if field in ("LimViolLimit", "LimViolValue", "LimViolPct"):
                        try:
                            value = float(str(raw).replace("%", ""))
                        except (ValueError, TypeError):
                            value = raw
                    else:
                        value = raw
                    cell = ws.cell(row_number, col, value)
                    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if isinstance(value, float):
                        cell.number_format = "0.0"
            if not rows:
                ws["B4"] = "No matching P1 violations"
            for column, width in {"B": 52, "C": 52, "D": 15, "E": 21, "F": 20, "G": 19}.items():
                ws.column_dimensions[column].width = width
            ws.freeze_panes = "D4"
            if rows:
                ws.auto_filter.ref = f"B3:G{len(rows) + 3}"
        if not book.sheetnames:
            continue
        path = company_dir / f"{name} P1 Results.xlsx"
        book.save(path)
        written.append(path)
        log(f"Wrote workbook: {path}")
    return written


def run_neighbor_studies(root: Path, assets_dir: Path, log, simauto_factory=None) -> tuple[list[Path], list[str]]:
    cases, warnings = discover_n1_cases(root)
    if not cases:
        raise ValueError("No N-1 source cases were found in the study folders.")
    results_root = root / RESULTS_DIR
    completed: dict[str, dict[str, list[dict[str, str]]]] = {}
    for case in cases:
        try:
            log(f"Starting study: {case.study}")
            completed[case.study] = run_n1_case(case, assets_dir, results_root, log, simauto_factory)
        except Exception as exc:
            warning = f"{case.study}: {exc}"
            warnings.append(warning)
            log(f"ERROR: {warning}")
    if not completed:
        raise RuntimeError("All N-1 studies failed. Review the log for individual errors.")
    workbooks = write_company_workbooks(results_root, completed, log)
    return workbooks, warnings
