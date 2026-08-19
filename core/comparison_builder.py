import os
import pandas as pd

from .case_types import CANONICAL_TO_PRETTY, TARGET_PATTERNS
from .column_blacklist import VOLTAGE_VIOLATION_CATEGORIES

# Try to import openpyxl for formatting + outline grouping
try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

def _to_float_series(series: pd.Series) -> pd.Series:
    """Convert a LimViolPct-like series to float safely."""
    if series is None:
        return pd.Series(dtype="float64")
    if pd.api.types.is_numeric_dtype(series):
        return series.astype(float)
    cleaned = series.astype(str).str.replace("%", "", regex=False).str.strip()
    return pd.to_numeric(cleaned, errors="coerce")


def _as_float(x):
    """Best-effort convert to float; returns None if not numeric."""
    try:
        if x is None:
            return None
        if isinstance(x, float) and pd.isna(x):
            return None
        # handle "123.4%" strings
        if isinstance(x, str):
            s = x.strip().replace("%", "")
            if s == "":
                return None
            return float(s)
        return float(x)
    except Exception:
        return None


def _round1_if_numeric(x):
    """Round numeric values to 1 decimal, otherwise return original."""
    v = _as_float(x)
    if v is None:
        return x
    return round(v, 1)


def _round_if_numeric(x, digits: int):
    v = _as_float(x)
    if v is None:
        return x
    return round(v, digits)


def _is_voltage_row(row) -> bool:
    cat = row.get("LimViolCat", "") if hasattr(row, "get") else getattr(row, "LimViolCat", "")
    return str(cat or "").strip() in VOLTAGE_VIOLATION_CATEGORIES


def _limit_or_value_for_display(row, field_name: str, report_type: str):
    digits = 2 if report_type == "voltage" or _is_voltage_row(row) else 1
    if hasattr(row, "get"):
        return _round_if_numeric(row.get(field_name, ""), digits)
    return _round_if_numeric(getattr(row, field_name, ""), digits)


def _pct_for_display(row, report_type: str):
    if report_type == "voltage":
        return None
    if hasattr(row, "get"):
        return _round1_if_numeric(row.get("LimViolPct", ""))
    return _round1_if_numeric(getattr(row, "LimViolPct", ""))


def _number_format_for_report(report_type: str) -> str:
    return "0.00" if report_type == "voltage" else "0.0"


def _filter_by_percent_threshold(df: pd.DataFrame, threshold: float) -> pd.DataFrame:
    if df is None or df.empty or "LimViolPct" not in df.columns:
        return df
    pct = _to_float_series(df["LimViolPct"])
    return df[pct.fillna(float("-inf")) >= float(threshold)].copy()


def _write_no_voltage_issues_row(ws, row: int, data_font, thin_border, left_align, center):
    ws.cell(row=row, column=2).value = "None"
    ws.cell(row=row, column=3).value = "No Voltage Issues"
    ws.cell(row=row, column=7).value = ""
    for col in range(2, 8):
        cell = ws.cell(row=row, column=col)
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = left_align if col in (2, 3) else center
    return row + 1


def _no_voltage_issues_record(case_type: str) -> dict:
    return {
        "CaseType": case_type,
        "CTGLabel": "None",
        "LimViolID": "No Voltage Issues",
        "LimViolLimit": None,
        "LimViolValue": None,
        "LimViolPct": None,
        "Notes": "",
    }


def _build_simple_workbook(
    root_folder,
    folder_to_case_csvs,
    log_func=None,
    threshold: float = 80.0,
    report_type: str = "thermal",
):
    """
    Fallback: simple one-sheet-per-scenario workbook, no fancy formatting,
    and no outline dropdown grouping.
    """
    if not folder_to_case_csvs:
        if log_func:
            log_func("No data to build combined workbook.")
        return None

    workbook_path = os.path.join(root_folder, "Combined_ViolationCTG_Comparison.xlsx")

    if log_func:
        log_func(f"\nBuilding SIMPLE combined workbook:\n  {workbook_path}")

    try:
        writer = pd.ExcelWriter(workbook_path)
    except Exception as e:
        if log_func:
            log_func(f"ERROR: Could not create ExcelWriter: {e}")
        return None

    try:
        with writer as xls_writer:
            for folder_name, case_map in folder_to_case_csvs.items():
                dfs = []
                for label in TARGET_PATTERNS:
                    csv_path = case_map.get(label)
                    if csv_path:
                        try:
                            df = pd.read_csv(csv_path)
                            df.insert(0, "CaseType", label)
                            if report_type == "thermal":
                                df = _filter_by_percent_threshold(df, threshold)
                            if "Notes" not in df.columns:
                                df["Notes"] = ""
                            if df.empty:
                                df = pd.DataFrame([_no_voltage_issues_record(label)])
                            dfs.append(df)
                            continue
                        except Exception as e:
                            if log_func:
                                log_func(f"  [{folder_name}] WARNING: Failed to read {csv_path}: {e}")
                    dfs.append(pd.DataFrame([_no_voltage_issues_record(label)]))

                if not dfs:
                    continue

                combined = pd.concat(dfs, ignore_index=True)
                sheet_name = (folder_name or "Sheet").strip()[:31] or "Sheet"
                combined.to_excel(xls_writer, sheet_name=sheet_name, index=False)

    except Exception as e:
        if log_func:
            log_func(f"ERROR while building simple workbook: {e}")
        return None

    if log_func:
        log_func("Simple combined workbook build complete.")
    return workbook_path


def build_workbook(
    root_folder,
    folder_to_case_csvs,
    group_details: bool = True,
    log_func=None,
    threshold: float = 80.0,
    report_type: str = "thermal",
):
    """
    Build a combined Excel workbook with one sheet per subfolder.

    group_details=True:
        - Within each CaseType block, group rows by LimViolID.
        - Show the highest LimViolPct row per LimViolID.
        - Collapse (hide) the other contingencies under an Excel outline dropdown.
        - Sort the groups so the WORST (highest percent loading) issues appear first.
    """
    if not folder_to_case_csvs:
        if log_func:
            log_func("No data to build combined workbook.")
        return None

    if not OPENPYXL_AVAILABLE:
        if log_func:
            log_func("openpyxl not available; building simple combined workbook without special formatting.")
        return _build_simple_workbook(
            root_folder,
            folder_to_case_csvs,
            log_func,
            threshold=threshold,
            report_type=report_type,
        )

    # ---------------------------
    # Build scenario DataFrames
    # ---------------------------
    scenario_data = {}  # folder_name -> combined DataFrame

    for folder_name, case_map in folder_to_case_csvs.items():
        dfs = []
        for label in TARGET_PATTERNS:
            csv_path = case_map.get(label)
            if not csv_path:
                continue
            try:
                df = pd.read_csv(csv_path)
                df.insert(0, "CaseType", label)
                dfs.append(df)
            except Exception as e:
                if log_func:
                    log_func(f"  [{folder_name}] WARNING: Failed to read {csv_path}: {e}")

        if dfs:
            scenario_data[folder_name] = pd.concat(dfs, ignore_index=True)

    if not scenario_data:
        if log_func:
            log_func("No scenario data to write into workbook.")
        return None

    # ---------------------------
    # Create formatted workbook
    # ---------------------------
    workbook_path = os.path.join(root_folder, "Combined_ViolationCTG_Comparison.xlsx")

    if log_func:
        log_func(f"\nBuilding FORMATTED combined workbook:\n  {workbook_path}")
        log_func(f"Expandable dropdown grouping is {'ON' if group_details else 'OFF'}.")
        if report_type == "thermal":
            log_func(f"Percent loading threshold: {float(threshold):.2f}%")
        else:
            log_func("Voltage report selected; percent loading threshold is ignored.")
        if report_type == "thermal":
            log_func("Sorting Resulting Issues by highest Percent Loading (worst first).")
        else:
            log_func("Voltage report values are displayed in p.u. to two decimals.")
        log_func("Shifting output down by 1 row (blank Row 1).")

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    title_fill = PatternFill(fill_type="solid", fgColor="305496")
    title_font = Font(color="FFFFFF", bold=True, size=12)

    header_fill = PatternFill(fill_type="solid", fgColor="305496")
    header_font = Font(color="FFFFFF", bold=True)

    data_font = Font(color="000000")
    data_bold_font = Font(color="000000", bold=True)

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Your updated required columns
    required_cols = ["CTGLabel", "LimViolLimit", "LimViolValue", "LimViolPct"]

    for folder_name, df in scenario_data.items():
        sheet_name = (folder_name or "Sheet").strip()[:31] or "Sheet"
        ws = wb.create_sheet(title=sheet_name)

        # Excel outline behavior: summary rows ABOVE details (so dropdown is on the max row)
        ws.sheet_properties.outlinePr.summaryBelow = False

        # Set column widths – contiguous columns B to G
        ws.column_dimensions["B"].width = 55  # Contingency Events
        ws.column_dimensions["C"].width = 55  # Resulting Issue
        ws.column_dimensions["D"].width = 18  # Limit
        ws.column_dimensions["E"].width = 22  # Contingency Value
        ws.column_dimensions["F"].width = 18  # Percent Loading
        ws.column_dimensions["G"].width = 35  # Notes
        if report_type == "voltage":
            ws.column_dimensions["F"].hidden = True

        # Shift everything down by 1 row
        current_row = 2

        for label in TARGET_PATTERNS:
            block_df = df[df["CaseType"] == label].copy()
            pretty_name = CANONICAL_TO_PRETTY.get(label, label)
            if report_type == "thermal":
                block_df = _filter_by_percent_threshold(block_df, threshold)
            if "Notes" not in block_df.columns:
                block_df["Notes"] = ""
            value_header = "Contingency Value (p.u.)" if report_type == "voltage" else "Contingency Value (MVA)"

            # ===== Title row =====
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=7)
            c = ws.cell(row=current_row, column=2)
            c.value = pretty_name
            c.fill = title_fill
            c.font = title_font
            c.alignment = center
            for col in range(2, 8):  # B..G
                ws.cell(row=current_row, column=col).border = thin_border
            current_row += 1

            # ===== Header row =====
            headers = [
                ("B", "Contingency Events"),
                ("C", "Resulting Issue"),
                ("D", "Limit"),
                ("E", value_header),
                ("F", "Percent Loading"),
                ("G", "Notes"),
            ]
            for col_letter, text in headers:
                col_idx = ord(col_letter) - ord("A") + 1
                hc = ws.cell(row=current_row, column=col_idx)
                hc.value = text
                hc.fill = header_fill
                hc.font = header_font
                hc.alignment = center
                hc.border = thin_border
            current_row += 1

            # Validate columns
            for col in required_cols:
                if col not in block_df.columns and log_func:
                    log_func(f"  [{folder_name} / {label}] WARNING: column '{col}' missing.")

            has_limviolid = "LimViolID" in block_df.columns

            if block_df.empty:
                current_row = _write_no_voltage_issues_row(
                    ws,
                    current_row,
                    data_font,
                    thin_border,
                    left_align,
                    center,
                )
                current_row += 1
                continue

            if group_details and has_limviolid:
                # Numeric percent for sorting
                if "LimViolPct" in block_df.columns:
                    block_df["_pct_num"] = _to_float_series(block_df["LimViolPct"])
                else:
                    block_df["_pct_num"] = pd.Series([float("nan")] * len(block_df), index=block_df.index)

                # Sort within each LimViolID so max is first
                sort_cols = ["LimViolID", "_pct_num"]
                asc = [True, False]
                if "CTGLabel" in block_df.columns:
                    sort_cols.append("CTGLabel")
                    asc.append(True)

                block_df = block_df.sort_values(
                    by=sort_cols, ascending=asc, na_position="last", kind="mergesort"
                )

                # Order groups by their max % (worst first)
                group_max = (
                    block_df.groupby("LimViolID", dropna=False)["_pct_num"]
                    .max()
                    .reset_index()
                    .rename(columns={"_pct_num": "_group_max_pct"})
                )
                group_max["_lim_str"] = group_max["LimViolID"].astype(str)
                group_max = group_max.sort_values(
                    by=["_group_max_pct", "_lim_str"],
                    ascending=[False, True],
                    na_position="last",
                    kind="mergesort",
                )

                ordered_limviolid_values = list(group_max["LimViolID"])

                for limviolid in ordered_limviolid_values:
                    g = block_df[block_df["LimViolID"].eq(limviolid)].copy()
                    if g.empty:
                        continue

                    rows = list(g.itertuples(index=False))
                    if not rows:
                        continue

                    # Summary row (max)
                    r0 = rows[0]

                    ws.cell(row=current_row, column=2).value = getattr(r0, "CTGLabel", "")
                    ws.cell(row=current_row, column=3).value = getattr(r0, "LimViolID", "")

                    lim_cell = ws.cell(row=current_row, column=4)
                    val_cell = ws.cell(row=current_row, column=5)
                    pct_cell = ws.cell(row=current_row, column=6)
                    notes_cell = ws.cell(row=current_row, column=7)

                    lim_cell.value = _limit_or_value_for_display(r0, "LimViolLimit", report_type)
                    val_cell.value = _limit_or_value_for_display(r0, "LimViolValue", report_type)
                    pct_cell.value = _pct_for_display(r0, report_type)
                    notes_cell.value = getattr(r0, "Notes", "")

                    value_number_format = _number_format_for_report(report_type)
                    lim_cell.number_format = value_number_format
                    val_cell.number_format = value_number_format
                    pct_cell.number_format = "0.0"

                    for col in range(2, 8):
                        cell = ws.cell(row=current_row, column=col)
                        cell.font = data_bold_font
                        cell.border = thin_border
                        cell.alignment = left_align if col in (2, 3) else center

                    summary_row = current_row
                    current_row += 1

                    # Detail rows (collapsed)
                    detail_start = None
                    detail_end = None

                    for r in rows[1:]:
                        if detail_start is None:
                            detail_start = current_row

                        ws.cell(row=current_row, column=2).value = getattr(r, "CTGLabel", "")
                        ws.cell(row=current_row, column=3).value = ""

                        lim_cell = ws.cell(row=current_row, column=4)
                        val_cell = ws.cell(row=current_row, column=5)
                        pct_cell = ws.cell(row=current_row, column=6)
                        notes_cell = ws.cell(row=current_row, column=7)

                        lim_cell.value = _limit_or_value_for_display(r, "LimViolLimit", report_type)
                        val_cell.value = _limit_or_value_for_display(r, "LimViolValue", report_type)
                        pct_cell.value = _pct_for_display(r, report_type)
                        notes_cell.value = getattr(r, "Notes", "")

                        value_number_format = _number_format_for_report(report_type)
                        lim_cell.number_format = value_number_format
                        val_cell.number_format = value_number_format
                        pct_cell.number_format = "0.0"

                        for col in range(2, 8):
                            cell = ws.cell(row=current_row, column=col)
                            cell.font = data_font
                            cell.border = thin_border
                            cell.alignment = left_align if col in (2, 3) else center

                        detail_end = current_row
                        current_row += 1

                    if detail_start is not None and detail_end is not None:
                        ws.row_dimensions.group(
                            detail_start,
                            detail_end,
                            outline_level=1,
                            hidden=True,
                        )
                        ws.row_dimensions[summary_row].collapsed = True

            else:
                # No grouping: dump rows
                for _, row in block_df.iterrows():
                    ws.cell(row=current_row, column=2).value = row.get("CTGLabel", "")
                    ws.cell(row=current_row, column=3).value = row.get("LimViolID", "") if has_limviolid else ""

                    lim_cell = ws.cell(row=current_row, column=4)
                    val_cell = ws.cell(row=current_row, column=5)
                    pct_cell = ws.cell(row=current_row, column=6)
                    notes_cell = ws.cell(row=current_row, column=7)

                    lim_cell.value = _limit_or_value_for_display(row, "LimViolLimit", report_type)
                    val_cell.value = _limit_or_value_for_display(row, "LimViolValue", report_type)
                    pct_cell.value = _pct_for_display(row, report_type)
                    notes_cell.value = row.get("Notes", "")

                    value_number_format = _number_format_for_report(report_type)
                    lim_cell.number_format = value_number_format
                    val_cell.number_format = value_number_format
                    pct_cell.number_format = "0.0"

                    for col in range(2, 8):
                        cell = ws.cell(row=current_row, column=col)
                        cell.font = data_font
                        cell.border = thin_border
                        cell.alignment = left_align if col in (2, 3) else center

                    current_row += 1

            # One blank row between blocks
            current_row += 1

    try:
        wb.save(workbook_path)
    except Exception as e:
        if log_func:
            log_func(f"ERROR saving formatted workbook: {e}")
        return None

    if log_func:
        log_func("Formatted combined workbook build complete.")
    return workbook_path
